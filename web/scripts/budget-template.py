# Builds the envelope budget template that /tools/envelope-budget-template/ offers for download.
#
#   python scripts/budget-template.py        (from web/; needs openpyxl)
#
# Output: public/files/envelope-budget-template.xlsx  Budget, Transactions and How to use sheets.
#         The formulas use only SUM, SUMIFS and ROUND, so they work in Excel, Google Sheets and
#         Numbers.
#         public/files/envelope-budget-template.csv   the Transactions sheet's columns with two
#         sample rows. UrBudget's CSV import reads its Date, Payee, Memo, Outflow and Inflow
#         columns (desktop/src-tauri/src/bank_import.rs) and skips Category.
#
# Amounts on the Transactions sheet have no currency sign and dates are year-month-day, so a CSV
# saved from that sheet imports into UrBudget as it is.

import csv
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "public" / "files"
NAME = "envelope-budget-template"

# Rows 2 to 2,000 of Transactions feed the budget.
TX_LAST = 2000
# The budget table: the header row, then groups and categories, with room to add more below.
HEAD_ROW = 11
TABLE_LAST = 60

MONTH = "November 2026"
LEFT_OVER = 280.00
GROUPS = [
    # (group, [(category, carried over, assigned)])
    ("Bills", [
        ("Rent or mortgage", 0, 1500),
        ("Electricity", 0, 120),
        ("Phone and internet", 0, 110),
        ("Car insurance", 0, 140),
        ("Subscriptions", 0, 30),
    ]),
    ("Everyday", [
        ("Groceries", 35.20, 600),
        ("Gas and transit", 0, 180),
        ("Eating out", 20.00, 150),
        ("Household and personal", 0, 100),
        ("Fun money", 0, 100),
    ]),
    ("Savings goals", [
        ("Emergency fund", 1200.00, 300),
        ("Car repairs", 250.00, 100),
        ("Vacation", 400.00, 150),
        ("Gifts and holidays", 150.00, 100),
    ]),
]
TRANSACTIONS = [
    # (date, payee, category, memo, outflow, inflow)
    (date(2026, 11, 1), "Employer", "Income", "Paycheck", None, 1700.00),
    (date(2026, 11, 1), "Landlord", "Rent or mortgage", "November rent", 1500.00, None),
    (date(2026, 11, 3), "Grocery store", "Groceries", "Weekly shop", 85.40, None),
    (date(2026, 11, 5), "Electric company", "Electricity", "", 112.37, None),
    (date(2026, 11, 7), "Gas station", "Gas and transit", "", 42.10, None),
    (date(2026, 11, 8), "Pizza place", "Eating out", "", 27.50, None),
    (date(2026, 11, 10), "Grocery store", "Groceries", "Weekly shop", 96.15, None),
    (date(2026, 11, 12), "Grocery store", "Groceries", "Returned an item", None, 12.99),
    (date(2026, 11, 15), "Employer", "Income", "Paycheck", None, 1700.00),
    (date(2026, 11, 16), "Phone company", "Phone and internet", "", 110.00, None),
    (date(2026, 11, 18), "Auto repair shop", "Car repairs", "Oil change", 64.99, None),
]
# The CSV keeps two: income and a purchase, like the app's own CSV template.
CSV_ROWS = [TRANSACTIONS[0], TRANSACTIONS[2]]

HOW_TO = [
    ("title", "How to use this budget"),
    ("text", "This is an envelope budget: every dollar you have gets a job, a category, before you spend it. "
             "Type in the shaded cells. Everything else is worked out by formulas."),
    ("head", "Your first month"),
    ("step", "1. Set up your categories. On the Budget sheet, rename the sample groups and categories, delete the "
             "ones you don't need and add your own. For a new category, type its name in column A of an empty row "
             "in the table (or insert a row inside a group), then copy the Spent and Available formulas from the "
             "row above."),
    ("step", "2. Fill in Left over from last month. In your first month, that's the money in your checking and "
             "savings accounts today. Clear the sample amounts in Carried over."),
    ("step", "3. Record income. On the Transactions sheet, add each paycheck with the category Income and the "
             "amount in Inflow."),
    ("step", "4. Give every dollar a job. Type amounts in Assigned until Ready to assign is $0.00. Start with "
             "bills due before your next paycheck, then everyday spending, then savings goals."),
    ("step", "5. Record spending. Add each purchase on the Transactions sheet: the date, the payee, the category "
             "and the amount in Outflow. Type the category exactly as it's written on the Budget sheet. A refund "
             "goes in Inflow, with the category the purchase came from."),
    ("step", "6. Cover overspending. When a category's Available turns red, lower Assigned in a category that "
             "has money to spare and raise the overspent one by the same amount."),
    ("head", "Checks"),
    ("text", "Spending with no matching category, on the Budget sheet, should be $0.00. If it isn't, a "
             "transaction's Category is blank or doesn't match a category name exactly."),
    ("text", "List spending and income only. Moving money between your own accounts, like paying a credit card "
             "bill from checking, isn't spending: the purchases were recorded when you made them."),
    ("text", "The formulas read rows 2 to 2,000 of the Transactions sheet."),
    ("head", "Starting a new month"),
    ("step", "1. Save a copy of the file for the new month."),
    ("step", "2. In the copy, put each category's Available into its Carried over. Copy the Available column and "
             "paste it into Carried over as values only (Paste Special > Values in Excel, Paste special > Values "
             "only in Google Sheets, Paste Formula Results in Numbers)."),
    ("step", "3. Put the old Ready to assign into Left over from last month, and change the month."),
    ("step", "4. Set Assigned back to 0, or keep the amounts that repeat, and clear the old rows on the "
             "Transactions sheet."),
    ("head", "Excel, Google Sheets and Numbers"),
    ("text", "The formulas use only SUM, SUMIFS and ROUND, so the file works in all three. In Google Sheets, choose "
             "File > Import > Upload. In Numbers, open the file and it's converted."),
    ("head", "Moving to UrBudget"),
    ("text", "UrBudget does this without formulas, on Windows and macOS: rollover from month to month, targets, "
             "credit cards and imports of your bank's CSV, OFX, QFX and QIF files. Save the Transactions sheet as "
             "a CSV file and UrBudget's Import file reads it as it is: it uses Date, Payee, Memo, Outflow and "
             "Inflow and skips Category. Pick year-month-day as the date order."),
    ("link", "https://urbudget.app/tools/envelope-budget-template/"),
    ("text", "Free to use, copy and change."),
]

# Colors from the website: warm paper, a soft violet for groups and the total, pale yellow for typing.
INPUT = PatternFill("solid", fgColor="FFF6E0")
HEAD = PatternFill("solid", fgColor="F2F0EA")
GROUP = PatternFill("solid", fgColor="F0EBFF")
KEY = PatternFill("solid", fgColor="E4DCFF")
LINE = Side(style="thin", color="D2CCC0")
MONEY = '"$"#,##0.00;[Red]-"$"#,##0.00'
AMOUNT = "#,##0.00"
GREY = Font(color="67626F", italic=True)

TX = "Transactions!"
def tx(col):
    return f"{TX}${col}$2:${col}${TX_LAST}"


def budget_sheet(ws):
    ws.title = "Budget"
    ws.sheet_properties.tabColor = "6D4AE8"
    ws.column_dimensions["A"].width = 34
    for col in "BCDE":
        ws.column_dimensions[col].width = 15

    ws["A1"] = "Envelope budget"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = "Month"
    ws["B2"] = MONTH
    ws["B2"].fill = INPUT
    ws["A3"] = "Type in the shaded cells; the rest is worked out. Steps are on the How to use sheet."
    ws["A3"].font = GREY

    first, last = HEAD_ROW + 1, TABLE_LAST
    summary = [
        (4, "Left over from last month", LEFT_OVER,
         "Money you had but hadn't assigned. In your first month: what's in checking and savings."),
        (5, "Income this month", f'=SUMIFS({tx("F")},{tx("C")},"Income")',
         "Inflows on the Transactions sheet with the category Income."),
        (6, "Assigned this month", f"=SUM(C{first}:C{last})", ""),
        (7, "Ready to assign", "=ROUND(B4+B5-B6,2)", "Assign until this is $0.00: every dollar has a job."),
        (8, "Available in all categories", f"=SUM(E{first}:E{last})", ""),
        (9, "Spending with no matching category",
         f"=ROUND(SUM({tx('E')})-SUM({tx('F')})+B5-SUM(D{first}:D{last}),2)",
         "Should be $0.00. If not, a transaction's Category is blank or misspelled."),
    ]
    for row, label, value, note in summary:
        ws.cell(row, 1, label)
        cell = ws.cell(row, 2, value)
        cell.number_format = MONEY
        if note:
            ws.cell(row, 3, note).font = GREY
    ws["B4"].fill = INPUT
    for col in (1, 2):
        ws.cell(7, col).font = Font(bold=True, size=12)
        ws.cell(7, col).fill = KEY

    headers = ["Category", "Carried over", "Assigned", "Spent", "Available"]
    for col, text in enumerate(headers, 1):
        cell = ws.cell(HEAD_ROW, col, text)
        cell.font = Font(bold=True)
        cell.fill = HEAD
        cell.border = Border(bottom=LINE)
        if col > 1:
            cell.alignment = Alignment(horizontal="right")

    row = first
    for group, categories in GROUPS:
        for col in range(1, 6):
            ws.cell(row, col).fill = GROUP
        ws.cell(row, 1, group).font = Font(bold=True)
        row += 1
        for name, carried, assigned in categories:
            ws.cell(row, 1, name).alignment = Alignment(indent=1)
            ws.cell(row, 2, carried).fill = INPUT
            ws.cell(row, 3, assigned).fill = INPUT
            ws.cell(row, 4, f"=SUMIFS({tx('E')},{tx('C')},$A{row})-SUMIFS({tx('F')},{tx('C')},$A{row})")
            ws.cell(row, 5, f"=ROUND(B{row}+C{row}-D{row},2)")
            row += 1
    for r in range(first, last + 1):
        for col in range(2, 6):
            ws.cell(r, col).number_format = MONEY
    ws.cell(HEAD_ROW - 1, 1, "To add a category, type its name in an empty row of the table (down to row "
            f"{TABLE_LAST}) and copy Spent and Available from the row above.").font = GREY


def transactions_sheet(ws):
    ws.title = "Transactions"
    widths = {"A": 13, "B": 24, "C": 24, "D": 28, "E": 13, "F": 13}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    headers = ["Date", "Payee", "Category", "Memo", "Outflow", "Inflow"]
    for col, text in enumerate(headers, 1):
        cell = ws.cell(1, col, text)
        cell.font = Font(bold=True)
        cell.fill = HEAD
        cell.border = Border(bottom=LINE)
        if col > 4:
            cell.alignment = Alignment(horizontal="right")
    for r, values in enumerate(TRANSACTIONS, 2):
        for col, value in enumerate(values, 1):
            if value not in (None, ""):
                ws.cell(r, col, value)
    # Formats go on the columns and the sample cells, not on empty rows: a spreadsheet saves every
    # formatted row to CSV, and empty rows there would stop UrBudget's import.
    for col, fmt in (("A", "yyyy-mm-dd"), ("E", AMOUNT), ("F", AMOUNT)):
        ws.column_dimensions[col].number_format = fmt
        for r in range(2, len(TRANSACTIONS) + 2):
            ws[f"{col}{r}"].number_format = fmt
    # An Excel table grows, with its formats and filter buttons, as rows are typed below it.
    table = Table(displayName="TransactionList", ref=f"A1:F{len(TRANSACTIONS) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"


def how_to_sheet(ws):
    ws.title = "How to use"
    ws.column_dimensions["A"].width = 100
    for r, (kind, text) in enumerate(HOW_TO, 1):
        cell = ws.cell(r, 1, text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "title":
            cell.font = Font(size=16, bold=True)
        elif kind == "link":
            cell.hyperlink = text
            cell.font = Font(color="5B38D6", underline="single")
        elif kind == "head":
            cell.font = Font(size=12, bold=True)
            ws.row_dimensions[r].height = 26
            cell.alignment = Alignment(vertical="bottom")


def workbook():
    wb = Workbook()
    budget_sheet(wb.active)
    transactions_sheet(wb.create_sheet())
    how_to_sheet(wb.create_sheet())
    wb.active = 0
    # No values are saved with the formulas, so spreadsheet apps work them out when the file opens.
    wb.calculation.fullCalcOnLoad = True
    stamp = datetime(2026, 10, 17)
    wb.properties.creator = "UrBudget"
    wb.properties.title = "Envelope budget template"
    wb.properties.created = stamp
    wb.properties.modified = stamp
    return wb


def amount(value):
    return "" if value is None else f"{value:.2f}"


def write_csv(path):
    with path.open("w", newline="", encoding="utf-8") as f:
        out = csv.writer(f, lineterminator="\r\n")
        out.writerow(["Date", "Payee", "Category", "Memo", "Outflow", "Inflow"])
        for day, payee, category, memo, outflow, inflow in CSV_ROWS:
            out.writerow([day.isoformat(), payee, category, memo, amount(outflow), amount(inflow)])


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    xlsx = OUT / f"{NAME}.xlsx"
    workbook().save(xlsx)
    write_csv(OUT / f"{NAME}.csv")
    for path in (xlsx, OUT / f"{NAME}.csv"):
        print(f"{path.relative_to(ROOT)}  {path.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
